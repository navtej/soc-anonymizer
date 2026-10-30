"""Regression tests for soc_anonymize.py - synthetic data only, rules mode (no NER download).

Run:  uv run --with pytest --with rich --with openpyxl --with python-calamine pytest -q tests
"""
import csv
import importlib.util
import json
import os
import signal
import subprocess
import sys
import time
from pathlib import Path

import pytest

import fixtures

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "soc_anonymize.py"
ENV = dict(os.environ, ANON_KEY="test-key")


def load_module():
    spec = importlib.util.spec_from_file_location("soc_anonymize", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.KEY = b"test-key"
    return mod


def run(args, cwd, check=True):
    r = subprocess.run([sys.executable, str(SCRIPT), *map(str, args), "--no-progress"],
                       cwd=cwd, env=ENV, capture_output=True, text=True, timeout=300)
    if check:
        assert r.returncode == 0, r.stdout + r.stderr
    return r


def read_csv(p):
    return list(csv.reader(open(p, newline="", encoding="utf-8")))


def token_index(mapping):
    """original (lowercased) -> set of tokens"""
    inv = {}
    for tok, orig in mapping.items():
        inv.setdefault(orig.lower(), set()).add(tok)
    return inv


def assert_mapping_sane(mapping):
    inv = token_index(mapping)
    multi = {k: v for k, v in inv.items() if len(v) > 1}
    assert not multi, f"same original got several tokens: {multi}"
    double = [k for k, v in mapping.items() if v in mapping]
    assert not double, f"tokens were anonymized again: {double}"


def all_text(folder):
    out = []
    for p in Path(folder).rglob("*"):
        if p.is_file() and p.suffix in (".csv", ".json", ".ndjson", ".log", ".txt") and p.name != "map.json":
            out.append(p.read_text(encoding="utf-8", errors="ignore"))
        if p.suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(p)
            out += [str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value]
    return "\n".join(out)


# --------------------------------------------------------------------------- Jira

def test_jira_identity_columns_and_free_text(tmp_path):
    src = fixtures.build_jira(tmp_path)
    run([src.name, "-o", "out", "--no-ner", "--map", "out/map.json"], tmp_path)
    out = read_csv(tmp_path / "out" / src.name)
    H, row = out[0], out[1]
    m = json.load(open(tmp_path / "out" / "map.json"))
    assert_mapping_sane(m)

    assert row[H.index("Issue key")].startswith("TKT-")
    assert row[H.index("Issue id")].startswith("ISSUEID_")
    assert row[H.index("Assignee")].startswith("USER_")
    assert row[H.index("Assignee Id")].startswith("USERID_")
    assert row[H.index("Custom field (Tenant Name)")].startswith("TENANT_")

    user, tenant, key = (row[H.index(c)] for c in ("Assignee", "Custom field (Tenant Name)", "Issue key"))
    assert user in row[H.index("Summary")] and tenant in row[H.index("Summary")]
    assert user in row[H.index("Custom field (Description.)")]
    assert key in row[H.index("Custom field (Triage Evidence(QC))")]

    text = open(tmp_path / "out" / src.name, encoding="utf-8").read().lower()
    left = [s for s in fixtures.JIRA_SECRETS if s.lower() in text]
    assert not left, f"leaked: {left}"


def test_jira_non_identity_columns_unchanged(tmp_path):
    src = fixtures.build_jira(tmp_path)
    run([src.name, "-o", "out", "--no-ner"], tmp_path)
    a, b = read_csv(src), read_csv(tmp_path / "out" / src.name)
    for col in ["Status", "Custom field (Severity)", "Created", "Custom field (Risk Score)",
                "Custom field (Threat Category)", "Issue Type"]:
        i = a[0].index(col)
        assert [r[i] for r in a] == [r[i] for r in b], col


# --------------------------------------------------------------------------- Sentinel

def test_sentinel_columns_json_and_consistency(tmp_path):
    src = fixtures.build_sentinel(tmp_path)
    run([src.name, "-o", "out", "--no-ner", "--map", "out/map.json"], tmp_path)
    out = read_csv(tmp_path / "out" / src.name)
    H, row = out[0], out[1]
    m = json.load(open(tmp_path / "out" / "map.json"))
    assert_mapping_sane(m)

    expect = {"TenantId": "TENANTID_", "ProviderName": "PROVIDER_", "VendorName": "VENDOR_",
              "VendorOriginalId": "VENDORID_", "SourceComputerId": "HOSTID_", "SourceSystem": "SYSTEM_",
              "WorkspaceSubscriptionId": "SUBSCRIPTIONID_", "WorkspaceResourceGroup": "RG_"}
    for col, prefix in expect.items():
        assert row[H.index(col)].startswith(prefix), (col, row[H.index(col)])
    # alert title must survive
    assert row[H.index("DisplayName")] == "Suspicious inbox forwarding rule"
    assert row[H.index("AlertName")] == "Suspicious inbox forwarding rule"
    # tenant GUID glued to '_' in AlertType
    assert row[H.index("AlertType")].startswith(row[H.index("TenantId")] + "_")
    # same RG / subscription token inside paths and links
    for col in ("ResourceId", "AlertLink", "ExtendedLinks"):
        assert row[H.index("WorkspaceResourceGroup")] in row[H.index(col)], col
        assert row[H.index("WorkspaceSubscriptionId")] in row[H.index(col)], col
    # entity JSON: split UPN, hosts, mailbox
    ent = {e["Type"]: e for e in json.loads(row[H.index("Entities")])}
    assert ent["account"]["Name"].startswith("USER_") and ent["account"]["UPNSuffix"].startswith("DOMAIN_")
    assert ent["host"]["HostName"].startswith("HOST_") and ent["host"]["OMSAgentID"].startswith("HOSTID_")
    email_tok = row[H.index("CompromisedEntity")]
    assert ent["mailbox"]["MailboxPrimaryAddress"] == email_tok == ent["mailMessage"]["Recipient"]
    ext = json.loads(row[H.index("ExtendedProperties")])
    assert not any("@acmeholdings" in k for k in ext), "e-mail used as a JSON key leaked"

    text = open(tmp_path / "out" / src.name, encoding="utf-8").read().lower()
    left = [s for s in fixtures.SENTINEL_SECRETS if s.lower() in text]
    assert not left, f"leaked: {left}"


# --------------------------------------------------------------------------- multi-file

def test_cross_file_one_token_per_value(tmp_path):
    d = fixtures.build_multi(tmp_path / "multi")
    run(["multi", "-o", "out", "--no-ner", "--id-columns", "Ticket=ISSUEKEY", "--map", "out/map.json"], tmp_path)
    m = json.load(open(tmp_path / "out" / "map.json"))
    assert_mapping_sane(m)
    inv = token_index(m)
    for orig in ["priya sharma", "acme holdings pvt ltd", "globex corporation", "initech inc", "mssoc-1203"]:
        assert len(inv[orig]) == 1, orig
    text = all_text(tmp_path / "out")
    left = [s for s in ["Priya", "Sharma", "Acme", "Globex", "GLOBEX", "Initech", "Desai", "Verma", "MSSOC",
                        "mssoc", "Turner"] if s in text]
    assert not left, f"leaked: {left}"


def test_weak_variants_only_when_name_cased(tmp_path):
    fixtures.build_multi(tmp_path / "multi")
    run(["multi/weak.csv", "-o", "out", "--no-ner"], tmp_path)
    s = (tmp_path / "out" / "weak.csv").read_text()
    assert "global admin role reviewed" in s and "will escalate" in s
    assert "Global Bank" not in s and "Will checked" not in s and "GLOBAL logs" not in s


# --------------------------------------------------------------------------- units

@pytest.mark.parametrize("text", [
    "plain priya.sharma@acmeholdings.com end", "Priya.Sharma@AcmeHoldings.COM",
    "url ?u=priya.sharma%40acmeholdings.com&x=1", "defanged priya.sharma[@]acmeholdings[.]com",
    "spoken priya.sharma (at) acmeholdings (dot) com", "SMTP:Priya.Sharma@acmeholdings.com;",
    "html priya.sharma&#64;acmeholdings.com", r"json \npriya.sharma@acmeholdings.com",
    r"json \u003cpriya.sharma@acmeholdings.com\u003e", r"escaped priya.sharma\u0040acmeholdings.com"])
def test_email_forms_share_one_token(text):
    sa = load_module()
    an = sa.Anonymizer(sa.build_rules([]))
    assert sa.pseudo("EMAIL", "priya.sharma@acmeholdings.com") in an.scrub_str(text)


@pytest.mark.parametrize("text", ["see section 4.2 (at) page 7", "version 1.2 at build 3.4"])
def test_email_lookalikes_untouched(text):
    sa = load_module()
    assert sa.Anonymizer(sa.build_rules([])).scrub_str(text) == text


@pytest.mark.parametrize("header,cat", [
    ("Issue key", "ISSUEKEY"), ("Issue id", "ISSUEID"), ("Assignee", "USER"), ("Assignee Id", "USERID"),
    ("Custom field (Tenant Name)", "TENANT"), ("Summary", None), ("Status", None),
    ("Custom field (Analyst Verdict)", None), ("Reporter", "USER"), ("Account Name", "USER"),
    ("userPrincipalName", "EMAIL"), ("Computer", "HOST"), ("TenantId", "TENANTID"),
    ("WorkspaceSubscriptionId", "SUBSCRIPTIONID"), ("WorkspaceResourceGroup", "RG"),
    ("SourceComputerId", "HOSTID"), ("VendorOriginalId", "VENDORID"), ("ProviderName", "PROVIDER"),
    ("VendorName", "VENDOR"), ("SourceSystem", "SYSTEM"), ("CompromisedEntity", "ENTITY"),
    ("DisplayName", None), ("AlertName", None), ("SystemAlertId", None), ("UserDisplayName", "USER"),
    ("account_name", "TENANT"), ("AccountName", "USER"), ("recipient_address", "EMAIL"),
    ("senderDisplayName", "USER"), ("agentComputerName", "HOST"), ("agentDomain", "DOMAIN"),
    ("agentLastLoggedInUserName", "USER"), ("original_raw_payload", None), ("alert_json", None)])
def test_column_classification(header, cat):
    assert load_module().column_category(header) == cat


def test_delimiter_detection():
    import io
    sa = load_module()
    for delim, text in [(",", "a; b; c"), (";", "a, b, c"), (",", "x\ty | z")]:
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=delim)
        w.writerow(["Issue key", "Assignee", "Description"])
        for i in range(40):
            w.writerow([f"K-{i}", "Someone", f"{text}\nmore {text}"])
        assert sa.detect_delimiter(buf.getvalue(), ".csv") == delim


def test_chunks_are_lossless_and_bounded():
    sa = load_module()
    for text in ["The analyst reviewed alert %d. " * 700, "QUJD" * 3000, "a\n" + "x" * 2500 + "\nb"]:
        parts = sa.chunks(text)
        assert "".join(parts) == text and max(map(len, parts)) <= 1000


# --------------------------------------------------------------------------- NER interplay

STUB_NER = r'''
import importlib.util, re, sys
spec = importlib.util.spec_from_file_location("sa", sys.argv[1]); sa = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sa)
class StubModel:
    def extract_entities(self, text, labels, threshold, include_confidence, max_len):
        found = [{"text": m, "confidence": 0.95} for m in
                 re.findall(r"\b[A-Z][a-z]+ [A-Z][a-z]+\b|\b[A-Z]+_[0-9a-f]{8}\b|TKT-\d{7}", text)]
        return {"entities": {"person": found}}
class StubNER(sa.NER):
    def __init__(self, *a):
        self.hits = 0; self.model = StubModel(); self.threshold, self.name_threshold = 0.4, 0.6
sa.NER = StubNER
sys.argv = ["soc_anonymize.py"] + sys.argv[2:]
sa.main()
'''


def test_ner_never_retokenizes_tokens(tmp_path):
    fixtures.build_sentinel(tmp_path)
    fixtures.build_jira(tmp_path)
    (tmp_path / "stub.py").write_text(STUB_NER)
    r = subprocess.run([sys.executable, "stub.py", str(SCRIPT), "SecurityAlert.csv", "jira_export.csv",
                        "-o", "out", "--map", "out/map.json", "--no-progress"],
                       cwd=tmp_path, env=ENV, capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stdout + r.stderr
    assert_mapping_sane(json.load(open(tmp_path / "out" / "map.json")))


# --------------------------------------------------------------------------- CLI behaviour

def test_outdir_inside_input_refused(tmp_path):
    fixtures.build_multi(tmp_path / "multi")
    r = run(["multi", "-o", "multi/anon", "--no-ner"], tmp_path, check=False)
    assert r.returncode != 0 and "inside the input folder" in (r.stdout + r.stderr)
    assert not (tmp_path / "multi" / "anon").exists()


def test_folder_recursion_mirrors_structure(tmp_path):
    d = tmp_path / "exports"
    (d / "sub" / "deeper").mkdir(parents=True)
    (d / ".git").mkdir()
    fixtures.write_csv(d / "top.csv", ["Assignee", "Note"], [["Priya Sharma", "x"]])
    fixtures.write_csv(d / "sub" / "a.csv", ["Assignee", "Note"], [["Rahul Verma", "y"]])
    (d / "sub" / "deeper" / "b.json").write_text('{"value":[{"userPrincipalName":"bob@contoso.com"}]}')
    (d / ".git" / "x.csv").write_text("a\n1\n")
    (d / "sample_manifest.json").write_text("{}")
    run(["exports", "-o", "out", "--no-ner"], tmp_path)
    got = sorted(str(p.relative_to(tmp_path / "out")) for p in (tmp_path / "out").rglob("*") if p.is_file())
    assert got == ["sub/a.csv", "sub/deeper/b.json", "top.csv"]


@pytest.mark.skipif(os.name != "posix", reason="SIGINT semantics")
def test_ctrl_c_is_immediate(tmp_path):
    fixtures.write_csv(tmp_path / "rows.csv", ["Assignee", "Description"],
                       [[f"user{i}", f"Long description text number {i} about an incident"] for i in range(5)])
    slow = STUB_NER.replace(
        "class StubModel:",
        "import hashlib\nclass StubModel:").replace(
        "        found = [",
        "        hashlib.pbkdf2_hmac('sha256', b'x', b'y', 30_000_000)  # stuck in C like a torch call\n"
        "        found = [")
    (tmp_path / "slow.py").write_text(slow)
    p = subprocess.Popen([sys.executable, "slow.py", str(SCRIPT), "rows.csv", "-o", "out", "--no-progress"],
                         cwd=tmp_path, env=ENV, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    time.sleep(4)
    t0 = time.monotonic()
    p.send_signal(signal.SIGINT)
    out, _ = p.communicate(timeout=30)
    assert p.returncode == 130, out
    assert time.monotonic() - t0 < 1.0, "Ctrl-C took too long"
    assert not (tmp_path / "out" / "rows.csv").exists(), "partial output left behind"


def test_xlsx_control_characters(tmp_path):
    """The streamed path decodes _x0000_ to a raw control char openpyxl can't write."""
    src = fixtures.build_control_chars_xlsx(tmp_path)
    run([src.name, "-o", "out", "--no-ner", "--xlsx-full-max-mb", "0"], tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "out" / "ctrl.xlsx").active
    assert ws.max_row == 11
    assert "\ufffd" in ws["B2"].value


# --------------------------------------------------------------------------- SentinelOne

@pytest.mark.parametrize("mode", [[], ["--xlsx-full-max-mb", "0"]], ids=["full", "streamed"])
def test_sentinelone_xlsx_tenant_emails_and_cut_json(tmp_path, mode):
    """account_name is the customer; alert_json / original_raw_payload (raw mails, display
    names, JSON cut off at the cell limit) and recipient_address are scrubbed consistently."""
    fixtures.build_sentinelone_xlsx(tmp_path / "s1")
    run(["s1", "-o", "out", "--no-ner", "--map", "out/map.json", *mode], tmp_path)
    text = all_text(tmp_path / "out")
    leaks = [s for s in fixtures.S1_SECRETS if s.lower() in text.lower()]
    assert not leaks, leaks
    from openpyxl import load_workbook
    rows = list(load_workbook(tmp_path / "out" / "s1_alerts.xlsx").active.iter_rows(values_only=True))
    sa = load_module()
    tenant = sa.pseudo("TENANT", "Acme Holdings")
    assert {r[2] for r in rows[1:]} == {tenant}
    assert rows[1][1] == "High" and "Suspicious e-mail link clicked" in rows[1][4]
    for r in rows[1:3]:                                 # parsed and cut-off JSON: same tenant token
        assert f'"accountName":"{tenant}"' in r[4].replace(" ", "")
        assert f'"accountName":"{tenant}"' in r[5].replace(" ", "")
    assert rows[1][3] == "; ".join(sa.pseudo("EMAIL", e) for e in
                                   ("priya.sharma@acmeholdings.com", "rahul.verma@acmeholdings.com"))
    assert sa.pseudo("IPV4", "10.20.5.17") in rows[1][4] and "mx.example.net" in rows[1][5]  # relay kept
    assert_mapping_sane(json.loads((tmp_path / "out" / "map.json").read_text()))


# --------------------------------------------------------------------------- Wazuh / AWS

def test_wazuh_fields_hosts_with_ips_and_aws(tmp_path):
    fixtures.build_wazuh(tmp_path / "wz")
    run(["wz", "-o", "out", "--no-ner", "--map", "out/map.json"], tmp_path)
    m = json.load(open(tmp_path / "out" / "map.json"))
    assert_mapping_sane(m)
    text = all_text(tmp_path / "out")
    left = [s for s in fixtures.WZ_SECRETS if s.lower() in text.lower()]
    assert not left, f"leaked: {left}"

    H, row = (r := read_csv(tmp_path / "out" / "wazuh.csv"))[0], r[1]
    col = lambda name: row[H.index(name)]
    # host names that embed an IP: one HOST token, in its column and in full_log
    host = col("_source.agent.name")
    assert host.startswith("HOST_") and col("_source.predecoder.hostname") == host
    assert host in col("_source.full_log")
    assert col("_source.manager.name").startswith("HOST_")
    # same IP -> same token in srcip, full_log and the dashed host name's own IP stays hidden
    assert col("_source.data.srcip") in col("_source.full_log")
    assert col("_source.data.aws.sourceIPAddress").startswith("HOST_")
    # public AWS service endpoints / principals kept, IP-bearing public DNS replaced
    assert col("_source.data.aws.eventSource") == "ec2.amazonaws.com"
    assert col("_source.data.aws.userIdentity.invokedBy") == "AWS Internal"
    assert col("_source.data.aws.resource.instanceDetails.networkInterfaces.publicDnsName").startswith("HOST_")
    # account id: one token in its column, the ARN and the S3 log path; ARN structure kept
    acct = col("_source.data.aws.userIdentity.accountId")
    assert acct.startswith("AWSACCOUNT_")
    assert col("_source.data.aws.userIdentity.arn").startswith(f"arn:aws:iam::{acct}:user/USER_")
    assert col("_source.data.aws.log_info.log_file").count(acct) == 2
    assert col("_source.data.srcuser") in col("_source.data.aws.userIdentity.arn")
    assert col("_source.data.aws.userIdentity.accessKeyId").startswith("AWSKEY_")
    assert col("_source.data.aws.responseElements.securityGroupRuleSet.items.securityGroupRuleId").startswith("AWSRES_")
    assert col("_source.data.aws.resource.instanceDetails.networkInterfaces.ipv6Addresses").startswith("2001:db8::")
    tags = json.loads(col("_source.data.aws.resource.instanceDetails.tags"))
    assert tags[0]["key"] == "Name" and tags[0]["value"].startswith("AWSTAG_")
    for g in ("city_name", "region_name", "country_name"):
        assert col(f"_source.GeoLocation.{g}").startswith("GEO_")
    # non-identity fields unchanged
    assert col("_source.rule.mail") == "false" and col("_source.rule.level") == "5"
    assert col("_source.rule.description") == "sshd: authentication success."
    assert col("_source.location") == "/var/log/secure"

    # nested JSON export: same tokens as the CSV, raw CloudTrail inside full_log scrubbed too
    hit = json.loads((tmp_path / "out" / "wazuh.ndjson").read_text())
    src = hit["_source"]
    assert src["manager"]["name"].startswith("HOST_") and src["rule"]["mail"] is False
    assert src["GeoLocation"]["city_name"] == col("_source.GeoLocation.city_name")
    aws = src["data"]["aws"]
    assert aws["userIdentity"]["accountId"] == acct == aws["recipientAccountId"]
    assert aws["log_info"]["s3bucket"] == col("_source.data.aws.log_info.s3bucket")
    assert aws["resource"]["instanceDetails"]["networkInterfaces"][0]["privateDnsName"] == host
    ct = json.loads(src["full_log"])
    assert ct["userIdentity"]["arn"] == aws["userIdentity"]["arn"]
    # role name is also sessionIssuer.userName -> its USER token, also inside the ARN
    role = ct["userIdentity"]["sessionContext"]["sessionIssuer"]["userName"]
    assert ct["userIdentity"]["arn"].startswith(f"arn:aws:sts::{acct}:assumed-role/{role}/USER_")
    assert ct["userIdentity"]["sessionContext"]["sessionIssuer"]["arn"] == f"arn:aws:iam::{acct}:role/{role}"
    assert ct["eventSource"] == "ec2.amazonaws.com" and ct["eventName"] == "AuthorizeSecurityGroupIngress"


@pytest.mark.parametrize("header,cat", [
    ("_source.agent.name", "HOST"), ("_source.manager.name", "HOST"), ("_source.predecoder.hostname", "HOST"),
    ("_source.data.srcuser", "USER"), ("_source.data.dstuser", "USER"), ("_source.rule.mail", None),
    ("_source.rule.description", None), ("_source.agent.id", None), ("_source.full_log", None),
    ("_source.data.win.eventdata.targetUserName", "USER"), ("_source.data.win.system.computer", "HOST"),
    ("_source.data.aws.userIdentity.accountId", "AWSACCOUNT"), ("_source.GeoLocation.city_name", "GEO"),
    ("_source.data.aws.userIdentity.sessionContext.sessionIssuer.principalId", "USERID"),
    ("_source.data.aws.resource.instanceDetails.networkInterfaces.0.privateIpAddress", "HOST"),
    ("_source.data.aws.eventSource", "HOST"), ("OwnerId", "USERID"), ("Custom field (Description.)", None)])
def test_dotted_field_classification(header, cat):
    assert load_module().column_category(header) == cat


@pytest.mark.parametrize("text,gone", [
    ("conn from ip-10-1-2-3 ok", "10-1-2-3"), ("dns ip-10-1-2-3.eu-west-1.compute.internal.", "10-1-2-3"),
    ("v6 2600:1f18:aa:bb::51 seen", "2600:1f18"), ("key AKIAZZEXAMPLE7QK3WPL used", "AKIAZZ"),
    ("arn:aws:s3:::acme-bucket/priv/x.txt", "acme-bucket")])
def test_aws_and_dashed_ip_rules(text, gone):
    sa = load_module()
    assert gone not in sa.Anonymizer(sa.build_rules([])).scrub_str(text)


@pytest.mark.parametrize("text", ["at 10:30:45 build 1.2.3", "arn:aws:iam::aws:policy/AdministratorAccess",
                                  "std::vector size 12:34", "event 4624 port 443"])
def test_aws_rules_leave_lookalikes(text):
    sa = load_module()
    assert sa.Anonymizer(sa.build_rules([])).scrub_str(text) == text


# --------------------------------------------------------------------------- streaming

def test_json_document_streamed_record_by_record(tmp_path):
    fixtures.build_streaming_set(tmp_path / "in")
    run(["in/azure.json", "in/array.json", "-o", "out", "--no-ner", "--map", "out/map.json"], tmp_path)
    assert_mapping_sane(json.load(open(tmp_path / "out" / "map.json")))
    for name, n in [("azure.json", 1500), ("array.json", 50)]:
        text = (tmp_path / "out" / name).read_text(encoding="utf-8")
        data = json.loads(text)
        assert text == json.dumps(data, ensure_ascii=False, indent=2), "layout differs from a whole dump"
        recs = data["value"] if name == "azure.json" else data
        assert len(recs) == n
        assert len({r["userPrincipalName"] for r in recs}) == 1          # one token for one original
        assert all(r["tags"] == [] and r["extra"] == {} and isinstance(r["ok"], bool) for r in recs)
        assert not any(s in text.lower() for s in ("priya", "acme", "10.20.30.")), name
    az = json.loads((tmp_path / "out" / "azure.json").read_text())
    assert list(az) == ["@odata.context", "count", "value", "nextLink"] and az["count"] == 1500
    assert az["nextLink"] is None


def test_ndjson_saved_as_json(tmp_path):
    fixtures.build_streaming_set(tmp_path / "in")
    run(["in/lines.json", "-o", "out", "--no-ner"], tmp_path)
    lines = (tmp_path / "out" / "lines.json").read_text().splitlines()
    assert len(lines) == 20 and all(json.loads(l)["userPrincipalName"].endswith("@example.com") for l in lines)


@pytest.mark.parametrize("script", ["soc_anonymize.py", "sample_files.py"])
def test_json_stream_parses_exactly_with_bounded_buffer(script):
    import io
    spec = importlib.util.spec_from_file_location(script[:-3], ROOT / script)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    tricky = [12345678901234567890, -0.5e-7, "quote \" and \\ and \u00e9 \ud83d\ude00", None, True,
              {"a": [1, {"b": []}], "": {}}, [], "x" * 50, 1e308, "\u2028 line sep"]
    doc = json.dumps({"head": 1, "value": tricky * 300, "tail": "end"}, ensure_ascii=False)
    for chunk in (1, 3, 7, 64):
        js = mod.JsonStream(io.StringIO(doc), chunk=chunk)
        got, peak, rest = [], 0, {}
        for k in js.keys():
            if js.peek() == "[":
                for _ in js.elements():
                    got.append(js.value())
                    peak = max(peak, len(js.buf))
            else:
                rest[k] = js.value()
        js.finish()
        assert got == tricky * 300 and rest == {"head": 1, "tail": "end"}, chunk
        assert peak <= 4 * max(chunk, 64), f"buffer grew to {peak} - not streaming"
    for bad in ('{"a": 1} {"b": 2}', '[1, 2', '[1 2]', '{"a" 1}'):     # extra data, cut off, broken
        js = mod.JsonStream(io.StringIO(bad), chunk=2)
        with pytest.raises(ValueError):
            if js.peek() == "[":
                for _ in js.elements():
                    js.value()
            else:
                for _ in js.keys():
                    js.value()
            js.finish()


def test_xlsx_streamed_path(tmp_path):
    """Large-book path: rows streamed, CRLF escapes decoded before the rules run,
    formatting-only rows below the data dropped, hidden sheets kept hidden."""
    fixtures.build_streaming_set(tmp_path / "in")
    r = run(["in/alerts.xlsx", "-o", "out", "--no-ner", "--xlsx-full-max-mb", "0"], tmp_path)
    assert "streamed" in r.stderr
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "out" / "alerts.xlsx")
    ws, raw = wb["Alerts"], wb["Raw"]
    assert ws.max_row == 31 and raw.sheet_state == "hidden"
    assert [c.value for c in ws[1]] == ["Account", "Notes"]
    cells = [str(c.value) for s in wb.worksheets for row in s.iter_rows() for c in row if c.value]
    assert not [c for c in cells if "10.1.2." in c or "Priya" in c or "192.168.7.7" in c], cells[:5]
    assert all("checked by SOC" in str(ws.cell(row=i, column=2).value) for i in range(2, 32))



def test_xlsx_stream_readers_agree(tmp_path):
    """calamine (default) and the openpyxl read-only fallback for huge books yield the
    same sheets and rows: leading empty rows/columns kept, chart sheets skipped."""
    src = fixtures.build_offset_xlsx(tmp_path)
    sa = load_module()

    def read():                                       # calamine pads rows to the sheet width
        return [(t, h, [r[:len(r) - next((i for i, v in enumerate(reversed(r)) if v is not None), len(r))]
                        for r in rows]) for t, h, rows in sa.xlsx_stream(src)]
    fast = read()
    sa.CALAMINE_MAX_MB = -1
    slow = read()
    assert fast == slow
    assert [(t, h) for t, h, _ in fast] == [("Offset", False), ("Hidden", True)]
    rows = fast[0][2]
    assert rows[:2] == [[], []] and rows[2][:3] == [None, None, "Assignee"]
    assert rows[4][2] == "Login from 10.9.8.7\rok" and rows[5] == [] and rows[6][2] == "after the gap"
    assert sa.xlsx_est_rows(src) == 7 + 2


def test_scrub_memo_keeps_column_context(tmp_path):
    """The scrub_str cache is keyed on the column too: the same value is a user in
    Assignee but free text in Summary."""
    sa = load_module()
    anon = sa.Anonymizer(sa.build_rules([]))
    as_user = anon.scrub_str("Priya Sharma", "Assignee")
    assert as_user == sa.pseudo("USER", "Priya Sharma")
    assert anon.scrub_str("Priya Sharma", "Summary") == "Priya Sharma"   # no lexicon, no NER
    assert anon.scrub_str("Priya Sharma", "Assignee") == as_user


def test_xlsx_crlf_escape_does_not_hide_ips_full_path(tmp_path):
    """Excel stores CR as _x000D_; '10.1.2.0_x000D_' has no word boundary, so the IP
    rule used to miss it in the full-fidelity path."""
    fixtures.build_streaming_set(tmp_path / "in")
    run(["in/alerts.xlsx", "-o", "out", "--no-ner"], tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "out" / "alerts.xlsx")["Alerts"]
    notes = [ws.cell(row=i, column=2).value for i in range(2, 32)]
    assert not [n for n in notes if "10.1.2." in n or "_x000D_" in n], notes[:3]


def test_failed_file_leaves_no_partial_output(tmp_path):
    good = "".join(json.dumps({"user": f"user{i}", "msg": "x" * 200}) + "\n" for i in range(2000))
    (tmp_path / "broken.ndjson").write_text(good + "{not json\n" + good)
    r = run(["broken.ndjson", "-o", "out", "--no-ner"], tmp_path, check=False)
    assert r.returncode == 1 and "broken.ndjson" in r.stderr
    assert not (tmp_path / "out" / "broken.ndjson").exists()


# --------------------------------------------------------------------------- parallel review, NER device

def test_ordered_map_keeps_order_and_runs_concurrently():
    sa = load_module()
    import threading
    live, peak, lock = [0], [0], threading.Lock()

    def slow(i):
        with lock:
            live[0] += 1; peak[0] = max(peak[0], live[0])
        time.sleep(0.02 * (i % 3))            # later items often finish first
        with lock:
            live[0] -= 1
        return i * i

    assert list(sa.ordered_map(slow, range(40), workers=8)) == [i * i for i in range(40)]
    assert peak[0] > 1
    assert list(sa.ordered_map(slow, range(5), workers=1)) == [0, 1, 4, 9, 16]

    def boom(i):
        if i == 7:
            raise ValueError("bad record")
        return i
    with pytest.raises(ValueError, match="bad record"):
        list(sa.ordered_map(boom, range(50), workers=4))


def _fake_reviewer(sa, threading):
    """A Reviewer without the HTTP server: it flags the invented codename 'Zorblax'
    after a short delay, and records how many requests were in flight at once."""
    rv = object.__new__(sa.Reviewer)
    rv.url, rv.apply, rv.min_len, rv.timeout, rv.model = "http://127.0.0.1:1/v1", True, 20, 5, "stub"
    rv.report, rv.cache, rv.calls, rv.flagged, rv.file = None, {}, 0, 0, ""
    rv.lock, rv.on_call = threading.Lock(), lambda: None
    rv.live, rv.peak, plock = 0, 0, threading.Lock()

    def ask(text):
        with plock:
            rv.live += 1; rv.peak = max(rv.peak, rv.live)
        time.sleep(0.01)
        with plock:
            rv.live -= 1
        return [{"text": "Zorblax", "type": "PROJECT"}] if "Zorblax" in text else []
    rv._ask = ask
    return rv


def test_parallel_review_output_matches_sequential(tmp_path):
    """--review-workers > 1 scrubs records in a thread pool: every format must come
    out byte-for-byte as in a sequential run, in the same order."""
    import threading
    from openpyxl import Workbook, load_workbook
    sa = load_module()
    src = tmp_path / "in"; src.mkdir()
    note = lambda i: f"Analyst {i} saw project Zorblax beaconing from 10.1.{i % 250}.7 during triage"
    fixtures.write_csv(src / "t.csv", ["Assignee", "Description"],
                       [[f"user{i}", note(i)] for i in range(60)])
    (src / "t.ndjson").write_text("".join(json.dumps({"user": f"u{i}", "msg": note(i)}) + "\n"
                                          for i in range(60)))
    (src / "t.json").write_text(json.dumps({"value": [{"Description": note(i)} for i in range(60)]}))
    (src / "t.log").write_text("".join(note(i) + "\n" for i in range(60)))
    wb = Workbook(); ws = wb.active; ws.title = "Alerts"
    ws.append(["Assignee", "Description"])
    for i in range(60):
        ws.append([f"user{i}", note(i)])
    wb.save(src / "t.xlsx")

    def run_with(workers, xlsx_full_max_mb):
        rv = _fake_reviewer(sa, threading)
        anon = sa.Anonymizer(sa.build_rules([]), reviewer=rv, workers=workers)
        out = tmp_path / f"out{workers}_{xlsx_full_max_mb}"
        with sa.UI(5, None, rv, enabled=False) as ui:
            for f in sorted(src.iterdir()):
                sa.process_file(f, out / f.name, anon, ui, xlsx_full_max_mb)
        return out, rv

    for mb in (40.0, 0.0):                         # XLSX full-fidelity and streamed paths
        seq, rv1 = run_with(1, mb)
        par, rv8 = run_with(8, mb)
        assert rv1.peak == 1 and rv8.peak > 1, (rv1.peak, rv8.peak)
        assert rv8.flagged == rv1.flagged > 0
        for f in sorted(seq.iterdir()):
            if f.suffix == ".xlsx":
                a = [r for r in load_workbook(f).active.iter_rows(values_only=True)]
                b = [r for r in load_workbook(par / f.name).active.iter_rows(values_only=True)]
                assert a == b and len(a) == 61 and not any("Zorblax" in r[1] for r in a)
            else:
                assert f.read_bytes() == (par / f.name).read_bytes(), f.name
                assert "Zorblax" not in f.read_text(), f.name
        rows = read_csv(par / "t.csv")
        assert [r[1].split()[1] for r in rows[1:]] == [str(i) for i in range(60)]   # order kept


def test_ner_device_selection(monkeypatch):
    import types
    sa = load_module()
    fake_torch = types.SimpleNamespace(cuda=types.SimpleNamespace(is_available=lambda: False),
                                       backends=types.SimpleNamespace(mps=types.SimpleNamespace(is_available=lambda: True)))
    monkeypatch.setitem(sys.modules, "torch", fake_torch)
    assert sa.pick_device("auto") == "mps"
    fake_torch.cuda.is_available = lambda: True
    assert sa.pick_device("auto") == "cuda"
    assert sa.pick_device("cpu") == "cpu"

    seen = {}

    class FakeGLiNER2:
        @classmethod
        def from_pretrained(cls, name, **kw):
            seen.update(kw, name=name)
            return cls()
    monkeypatch.setitem(sys.modules, "gliner2", types.SimpleNamespace(GLiNER2=FakeGLiNER2))
    sa.NER("some/model", 0.4, 0.6)                  # auto -> cuda with the fake torch
    assert seen == {"name": "some/model", "map_location": "cuda"}
    sa.NER("some/model", 0.4, 0.6, "cpu")
    assert seen["map_location"] == "cpu"
    fake_torch.cuda.is_available = lambda: False
    fake_torch.backends.mps.is_available = lambda: False
    assert sa.pick_device("auto") == "cpu"


# --------------------------------------------------------------------------- resume, progress log, outages

RESUME_ROWS = {
    # the name is an identity column only in a.csv; the others mention it only in prose
    "a.csv": (["Assignee", "Summary"], [["Priya Sharma", "Reviewed the Zorblax alert today"]]),
    "b.csv": (["Summary"], [["Escalated to Priya Sharma today"], ["Second Zorblax note for the team"]]),
    "c.csv": (["Summary"], [["Host check done by Priya Sharma"]]),
}


def resume_inputs(folder):
    folder.mkdir(parents=True, exist_ok=True)
    for name, (header, rows) in RESUME_ROWS.items():
        fixtures.write_csv(folder / name, header, rows)
    return folder


def read_log(path):
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines()]


def forget(manifest, *names):
    """State after an interruption: these files never finished."""
    m = json.loads(manifest.read_text())
    for n in names:
        del m["completed"][n]
    manifest.write_text(json.dumps(m))


def test_resume_skips_finished_keeps_lexicon_and_bytes(tmp_path):
    resume_inputs(tmp_path / "in")
    run(["in", "-o", "full", "--no-ner", "--map", "full.json"], tmp_path)
    run(["in", "-o", "out", "--no-ner", "--map", "out.json"], tmp_path)
    forget(tmp_path / "out.manifest.json", "b.csv", "c.csv")
    (tmp_path / "out" / "a.csv").write_bytes(b"DONE-MARKER\n")
    (tmp_path / "out" / "b.csv").unlink()
    (tmp_path / "out" / "c.csv").unlink()
    (tmp_path / "out" / ".part-c.csv").write_bytes(b"truncated by a hard kill")
    r = run(["in", "-o", "out", "--no-ner", "--map", "out.json", "--resume"], tmp_path)
    assert "already done" in r.stdout + r.stderr
    assert (tmp_path / "out" / "a.csv").read_bytes() == b"DONE-MARKER\n", "finished file was redone"
    assert b"Priya" not in (tmp_path / "out" / "b.csv").read_bytes(), "name from the skipped file got through"
    for name in ("b.csv", "c.csv"):
        assert (tmp_path / "out" / name).read_bytes() == (tmp_path / "full" / name).read_bytes()
    assert not (tmp_path / "out" / ".part-c.csv").exists()
    assert json.load(open(tmp_path / "out.json")) == json.load(open(tmp_path / "full.json"))
    # a second resume has nothing left to do, and outputs that vanished are done again
    (tmp_path / "out" / "c.csv").unlink()
    r = run(["in", "-o", "out", "--no-ner", "--resume"], tmp_path)
    assert "1 file(s) anonymized, 2 skipped" in r.stdout + r.stderr
    assert (tmp_path / "out" / "c.csv").read_bytes() == (tmp_path / "full" / "c.csv").read_bytes()


def test_resume_refuses_to_mix_runs(tmp_path):
    resume_inputs(tmp_path / "in")
    base = ["in", "-o", "out", "--no-ner"]
    run(base, tmp_path)
    before = (tmp_path / "out" / "a.csv").read_bytes()

    def refused(args, env=ENV):
        r = subprocess.run([sys.executable, str(SCRIPT), *args, "--resume", "--no-progress"],
                           cwd=tmp_path, env=env, capture_output=True, text=True, timeout=300)
        assert r.returncode != 0
        return r.stdout + r.stderr

    assert "ANON_KEY" in refused(base, dict(ENV, ANON_KEY="another-key"))
    assert "domains" in refused(base + ["--domains", "corp.local"])
    assert "no_lexicon" in refused(base + ["--no-lexicon"])
    fixtures.write_csv(tmp_path / "in" / "c.csv", ["Summary"], [["A longer, different export of this file"]])
    assert "input c.csv" in refused(base)
    (tmp_path / "in" / "c.csv").unlink()
    assert "input c.csv" in refused(base)
    resume_inputs(tmp_path / "in")
    (tmp_path / "out.manifest.json").unlink()
    assert "no record" in refused(base)
    assert (tmp_path / "out" / "a.csv").read_bytes() == before
    # nothing done yet: --resume is simply a first run, so retry loops can always pass it
    run(["in", "-o", "fresh", "--no-ner", "--resume"], tmp_path)
    assert (tmp_path / "fresh" / "a.csv").exists()


def test_fresh_run_does_not_inherit_finished_state(tmp_path):
    resume_inputs(tmp_path / "in")
    run(["in", "-o", "out", "--no-ner"], tmp_path)
    run(["in", "-o", "out", "--no-ner"], tmp_path)
    assert set(json.load(open(tmp_path / "out.manifest.json"))["completed"]) == set(RESUME_ROWS)
    # an interrupted second run must not vouch for the first run's outputs
    m = json.load(open(tmp_path / "out.manifest.json"))
    m["completed"] = {}
    (tmp_path / "out.manifest.json").write_text(json.dumps(m))
    r = run(["in", "-o", "out", "--no-ner", "--resume"], tmp_path)
    assert "3 file(s) anonymized" in r.stdout + r.stderr


def test_failed_file_keeps_no_partial_and_output_only_when_complete(tmp_path):
    sa = load_module()
    src = tmp_path / "rows.csv"
    fixtures.write_csv(src, ["Summary"], [["one note here"], ["second note here"]])
    dst = tmp_path / "out" / "rows.csv"

    class Boom(sa.Anonymizer):
        def scrub_str(self, *a, **k):
            assert not dst.exists(), "final name in use before the file is complete"
            assert sa.part_path(dst).exists()
            raise ValueError("boom")

    with sa.UI(1, enabled=False) as ui:
        with pytest.raises(ValueError):
            sa.process_file(src, dst, Boom(sa.build_rules([])), ui)
    assert not dst.exists() and not sa.part_path(dst).exists()


def test_xlsm_outputs_are_committed_under_the_right_name(tmp_path):
    import openpyxl
    (tmp_path / "in").mkdir()
    wb = openpyxl.Workbook()
    wb.active.append(["Assignee", "Note"])
    wb.active.append(["Priya Sharma", "x"])
    wb.save(tmp_path / "in" / "book.xlsm")
    run(["in", "-o", "full", "--no-ner"], tmp_path)
    assert [p.name for p in (tmp_path / "full").iterdir()] == ["book.xlsm"]
    run(["in", "-o", "streamed", "--no-ner", "--xlsx-full-max-mb", "0"], tmp_path)
    assert [p.name for p in (tmp_path / "streamed").iterdir()] == ["book.xlsx"]
    r = run(["in", "-o", "streamed", "--no-ner", "--xlsx-full-max-mb", "0", "--resume"], tmp_path)
    assert "already done" in r.stdout + r.stderr


def test_progress_log_records_the_run_without_values(tmp_path):
    resume_inputs(tmp_path / "in")
    (tmp_path / "in" / "old.xls").write_bytes(b"x")
    (tmp_path / "in" / "bad.json").write_text('{"value": [{"Summary": "SECRETVALUE"}')
    r = run(["in", "-o", "out", "--no-ner", "--domains", "acme.internal,corp.local",
             "--id-columns", "Hidden Header=USER"], tmp_path, check=False)
    assert r.returncode == 1            # bad.json
    text = (tmp_path / "out.progress.log").read_text()
    for secret in ("Priya", "SECRETVALUE", "acme.internal", "Hidden Header", "test-key"):
        assert secret not in text
    events = read_log(tmp_path / "out.progress.log")
    assert events[0]["event"] == "run_start" and events[0]["args"]["domains"] == 2
    assert events[-1]["event"] == "run_end" and events[-1]["exit_code"] == 1
    finished = {e["file"]: e for e in events if e["event"] == "file_finish"}
    assert finished["b.csv"]["records"] == 2 and finished["b.csv"]["seconds"] >= 0
    assert [e["file"] for e in events if e["event"] == "file_fail"] == ["bad.json"]
    assert [e["file"] for e in events if e["event"] == "file_skip"] == ["old.xls"]
    # appended, also to a chosen path; that path is never taken for an input
    run(["in/a.csv", "-o", "out", "--no-ner"], tmp_path)
    assert len(read_log(tmp_path / "out.progress.log")) > len(events)
    manifest = (tmp_path / "out.manifest.json").read_text()
    for secret in ("Priya", "acme.internal", "test-key"):
        assert secret not in manifest
    run(["in/a.csv", "-o", "o2", "--no-ner", "--progress-log", "logs/run.log"], tmp_path)
    assert read_log(tmp_path / "logs" / "run.log")[0]["event"] == "run_start"
    assert not (tmp_path / "o2.progress.log").exists()


# A review server the test controls: flags 'Zorblax', goes away after DOWN_AFTER requests.
STUB_REVIEW = r'''
import importlib.util, json, os, sys, threading, urllib.error
spec = importlib.util.spec_from_file_location("sa", sys.argv[1]); sa = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sa)
sa.RETRY_WAIT = (0.01, 0.01)
calls, lock = [0], threading.Lock()
class Reply:
    def __init__(self, content): self.body = json.dumps({"choices": [{"message": {"content": content}}]}).encode()
    def __enter__(self): return self
    def __exit__(self, *a): pass
    def read(self): return self.body
def urlopen(req, timeout=None):
    with lock:
        calls[0] += 1
        n = calls[0]
    if n > int(os.environ.get("DOWN_AFTER", 10**9)):
        raise urllib.error.URLError(ConnectionRefusedError(61, "Connection refused"))
    text = json.loads(req.data)["messages"][1]["content"]
    return Reply(json.dumps({"leaks": [{"text": "Zorblax", "type": "PROJECT"}] if "Zorblax" in text else []}))
sa.urllib.request.urlopen = urlopen
if os.environ.get("NO_CLOSE"):
    sa.close_write_only = lambda wb: None
sys.argv = ["soc_anonymize.py"] + sys.argv[2:]
sa.main()
'''


def review_run(tmp_path, args, down_after=None, **env):
    e = dict(ENV, **env)
    if down_after is not None:
        e["DOWN_AFTER"] = str(down_after)
    (tmp_path / "stub_review.py").write_text(STUB_REVIEW)
    return subprocess.run([sys.executable, "stub_review.py", str(SCRIPT), *map(str, args),
                           "--no-ner", "--no-progress", "--review", "--review-apply", "--review-model", "m",
                           "--review-retry-minutes", "0.002"],
                          cwd=tmp_path, env=e, capture_output=True, text=True, timeout=300)


@pytest.mark.parametrize("workers", ["1", "4"])
def test_review_outage_stops_cleanly_and_resume_continues(tmp_path, workers):
    resume_inputs(tmp_path / "in")
    extra = ["--review-workers", workers, "--map", "map.json"]
    full = review_run(tmp_path, ["in", "-o", "full", "--review-report", "full.jsonl", *extra])
    assert full.returncode == 0, full.stdout + full.stderr
    full_map = json.load(open(tmp_path / "map.json"))

    # a.csv needs 1 request (the identity column is short), b.csv dies on its second
    args = ["in", "-o", "out", "--review-report", "out.jsonl", *extra]
    r = review_run(tmp_path, args, down_after=2)
    assert r.returncode == 75, r.stdout + r.stderr
    assert "--resume" in r.stderr and "Traceback" not in r.stderr
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == ["a.csv"], "partial output left behind"
    events = read_log(tmp_path / "out.progress.log")
    assert events[-1]["event"] == "run_end" and events[-1]["reason"] == "review server unavailable"
    assert events[-1]["exit_code"] == 75
    assert {"event": "file_fail", "file": "b.csv"}.items() <= [e for e in events if e["event"] == "file_fail"][0].items()
    assert any(e["event"] == "review_outage" for e in events)
    flagged_before = (tmp_path / "out.jsonl").read_text()
    assert "a.csv" in flagged_before

    r = review_run(tmp_path, [*args, "--resume"])
    assert r.returncode == 0, r.stdout + r.stderr
    assert "already done" in r.stderr
    assert (tmp_path / "out.jsonl").read_text().startswith(flagged_before), "--resume wiped the report"
    for name in RESUME_ROWS:
        assert (tmp_path / "out" / name).read_bytes() == (tmp_path / "full" / name).read_bytes()
    assert json.load(open(tmp_path / "map.json")) == full_map


def test_streamed_xlsx_abort_closes_cleanly(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Summary"])
    for i in range(20):
        wb.active.append([f"Long enough note number {i} for the reviewer"])
    wb.create_sheet("Second").append(["Summary"])
    wb.save(tmp_path / "book.xlsx")
    args = ["book.xlsx", "-o", "out", "--xlsx-full-max-mb", "0", "--review-report", "r.jsonl"]
    r = review_run(tmp_path, args, down_after=5)
    assert r.returncode == 75, r.stdout + r.stderr
    assert "Exception ignored" not in r.stderr and "closed file" not in r.stderr, r.stderr
    assert list((tmp_path / "out").iterdir()) == []
    # the guard above sees the noise when the book is left to the garbage collector
    r = review_run(tmp_path, args, down_after=5, NO_CLOSE="1")
    assert r.returncode == 75 and "closed file" in r.stderr, r.stderr


def _reviewer(sa, monkeypatch, replies, minutes=1):
    """Each reply is raised, or returned as the response body."""
    sa.RETRY_WAIT = (0, 0)
    rv = sa.Reviewer("http://127.0.0.1:1/v1", model="m", retry_minutes=minutes)
    seen = []

    def urlopen(req, timeout=None):
        seen.append(req.full_url)
        reply = replies.pop(0)
        if isinstance(reply, BaseException):
            raise reply

        class R:
            def __enter__(self): return self
            def __exit__(self, *a): pass
            def read(self): return reply
        return R()

    monkeypatch.setattr(sa.urllib.request, "urlopen", urlopen)
    return rv, seen


def _http_error(sa, code):
    return sa.urllib.error.HTTPError("http://x", code, "err", {}, None)


def test_review_retries_connection_errors_and_timeouts(monkeypatch):
    sa = load_module()
    import http.client
    refused = sa.urllib.error.URLError(ConnectionRefusedError(61, "refused"))
    rv, seen = _reviewer(sa, monkeypatch, [refused, TimeoutError(), ConnectionResetError(), _http_error(sa, 503),
                                           http.client.IncompleteRead(b""), b"<html>starting", b'{"ok": 1}'])
    assert rv._http("/x") == {"ok": 1} and len(seen) == 7
    assert rv.down_since is None


@pytest.mark.parametrize("code", [400, 401, 404])
def test_review_4xx_is_not_retried(monkeypatch, code):
    sa = load_module()
    rv, seen = _reviewer(sa, monkeypatch, [_http_error(sa, code), b"{}"])
    with pytest.raises(sa.ReviewError) as e:
        rv._http("/x")
    assert len(seen) == 1 and e.value.exit_code == 1


def test_review_retries_are_bounded(monkeypatch):
    sa = load_module()
    refused = sa.urllib.error.URLError(ConnectionRefusedError(61, "refused"))
    rv, seen = _reviewer(sa, monkeypatch, [refused] * 10_000, minutes=0.001)
    sa.RETRY_WAIT = (0.01, 0.01)
    with pytest.raises(sa.ReviewError) as e:
        rv._http("/x")
    assert e.value.exit_code == sa.EXIT_RETRY_LATER and 1 <= len(seen) < 10_000
    assert rv.gave_up.is_set()
    n = len(seen)                       # the other workers stop at their next failure
    with pytest.raises(sa.ReviewError):
        rv._http("/x")
    assert len(seen) == n + 1


def test_review_startup_probe_fails_fast(monkeypatch):
    sa = load_module()
    calls = []

    def urlopen(req, timeout=None):
        calls.append(1)
        raise sa.urllib.error.URLError(ConnectionRefusedError(61, "refused"))

    monkeypatch.setattr(sa.urllib.request, "urlopen", urlopen)
    with pytest.raises(sa.ReviewError):
        sa.Reviewer("http://127.0.0.1:1/v1")
    assert len(calls) == 1


@pytest.mark.skipif(os.name != "posix", reason="SIGKILL semantics")
def test_hard_kill_leaves_no_output_under_the_final_name(tmp_path):
    fixtures.write_csv(tmp_path / "rows.csv", ["Assignee", "Description"],
                       [[f"user{i}", f"Long description text number {i} about an incident"] for i in range(5)])
    stub = STUB_NER.replace(
        "class StubModel:",
        "import hashlib, os\nclass StubModel:").replace(
        "        found = [",
        "        if os.environ.get('SLOW'): hashlib.pbkdf2_hmac('sha256', b'x', b'y', 30_000_000)\n"
        "        found = [")
    (tmp_path / "stub.py").write_text(stub)
    cmd = [sys.executable, "stub.py", str(SCRIPT), "rows.csv", "-o", "out", "--no-progress", "--resume"]
    p = subprocess.Popen(cmd, cwd=tmp_path, env=dict(ENV, SLOW="1"),
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    log_file = tmp_path / "out.progress.log"
    deadline = time.monotonic() + 60
    while "file_start" not in (log_file.read_text() if log_file.exists() else ""):
        assert time.monotonic() < deadline
        time.sleep(0.05)
    p.kill()
    p.wait()
    assert not (tmp_path / "out" / "rows.csv").exists()
    assert read_log(log_file)[-1]["event"] == "file_start", "a hard kill writes no run_end"
    r = subprocess.run(cmd, cwd=tmp_path, env=ENV, capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stdout + r.stderr
    assert (tmp_path / "out" / "rows.csv").exists()
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == ["rows.csv"]
    assert read_log(log_file)[-1]["event"] == "run_end"
