"""Regression tests for sample_files.py - synthetic data only."""
import hashlib
import json
import math
import os
import subprocess
import sys
from pathlib import Path

import fixtures

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "sample_files.py"


def run(args, cwd, check=True):
    r = subprocess.run([sys.executable, str(SCRIPT), *map(str, args), "--no-progress"],
                       cwd=cwd, capture_output=True, text=True, timeout=300)
    if check:
        assert r.returncode == 0, r.stdout + r.stderr
    return r


def digest(folder):
    return {str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(Path(folder).rglob("*")) if p.is_file()}


def test_exact_counts_and_formats(tmp_path):
    d = fixtures.build_sampling_set(tmp_path / "in")
    run(["in", "-o", "out", "--seed", "42"], tmp_path)
    m = json.load(open(tmp_path / "out" / "sample_manifest.json"))
    got = {Path(f["source"]).name: (f["kept"], f["total"]) for f in m["files"]}
    for name, total in [("jira.csv", 1000), ("semi.csv", 200), ("data.tsv", 500), ("siem.ndjson.gz", 2000),
                        ("azure.json", 300), ("elastic.json", 800), ("app.log", 100)]:
        assert got[name] == (math.ceil(total * 0.1), total), name
    assert got["book.xlsx"] == (40 + 4, 437)                     # per sheet: ceil(400*10%) + ceil(37*10%)
    assert "legacy.xls" not in got
    out = tmp_path / "out"
    assert (out / "siem.ndjson").exists() and (out / "sub" / "app.log").exists()
    az = json.load(open(out / "azure.json"))
    assert az["count"] == len(az["value"]) == 30
    assert (out / "semi.csv").read_text().splitlines()[0] == "a;b"
    assert (out / "jira.csv").read_bytes().startswith(b"\xef\xbb\xbf")   # BOM preserved
    from openpyxl import load_workbook
    wb = load_workbook(out / "book.xlsx")
    assert [(ws.title, ws.sheet_state) for ws in wb.worksheets] == [("Tickets", "visible"), ("Raw", "hidden")]


def test_reproducible_and_inputs_untouched(tmp_path):
    d = fixtures.build_sampling_set(tmp_path / "in")
    before = digest(d)
    run(["in", "-o", "a", "--seed", "7"], tmp_path)
    run(["in", "-o", "b", "--seed", "7"], tmp_path)
    run(["in", "-o", "c", "--seed", "8"], tmp_path)
    assert digest(d) == before
    for f in ["jira.csv", "siem.ndjson", "elastic.json", "sub/app.log"]:
        assert (tmp_path / "a" / f).read_bytes() == (tmp_path / "b" / f).read_bytes(), f
        assert (tmp_path / "a" / f).read_bytes() != (tmp_path / "c" / f).read_bytes(), f


def test_outdir_inside_input_refused(tmp_path):
    fixtures.build_sampling_set(tmp_path / "in")
    r = run(["in", "-o", "in/samples"], tmp_path, check=False)
    assert r.returncode != 0
    assert not (tmp_path / "in" / "samples").exists()


def test_line_endings_preserved(tmp_path):
    (tmp_path / "lf.csv").write_bytes(b"a,b\n1,x\n2,y\n")
    (tmp_path / "crlf.csv").write_bytes(b"a,b\r\n1,x\r\n2,y\r\n")
    run(["lf.csv", "crlf.csv", "-p", "100", "-o", "out"], tmp_path)
    assert (tmp_path / "out" / "lf.csv").read_bytes() == b"a,b\n1,x\n2,y\n"
    assert (tmp_path / "out" / "crlf.csv").read_bytes() == b"a,b\r\n1,x\r\n2,y\r\n"


def test_xlsx_control_characters(tmp_path):
    """calamine (fast path) decodes _x001E_ to a raw control char openpyxl can't write."""
    fixtures.build_control_chars_xlsx(tmp_path / "in")
    run(["in", "-o", "out", "-p", "100", "--xlsx-full-max-mb", "0"], tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "out" / "ctrl.xlsx").active
    assert ws.max_row == 11
    assert ws["A2"].value == "NT AU\ufffdTHORITY\\SYSTEM"
